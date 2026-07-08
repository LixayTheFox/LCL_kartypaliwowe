from datetime import date, datetime
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from archive_database import DatabaseConfig, load_database_config, save_database_config
from fuel_analysis import (
    SOURCE_TELEMATICS,
    UNKNOWN_VEHICLE,
    FuelRecord,
    analyze_records,
    export_xlsx,
    list_archive_reports,
    load_archive_report,
    load_single_report,
    normalize_plate,
    save_archive_report,
)


BASE_DIR = Path(__file__).resolve().parent
REPORT = BASE_DIR / "export_20260601085249.xlsx"


class FuelAnalysisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.records = load_single_report(REPORT)

    def test_normalizes_vehicle_registration(self):
        self.assertEqual(normalize_plate("BI 228HA"), "BI228HA")
        self.assertEqual(normalize_plate("bi-228 ha"), "BI228HA")

    def test_loads_single_report_format(self):
        self.assertEqual(len(self.records), 685)
        self.assertTrue(all(record.source == SOURCE_TELEMATICS for record in self.records))

    def test_loads_minimal_export_format_with_gasoline_and_missing_plate(self):
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "minimal.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "Numery rejestracyjne",
                    "Data",
                    "Ilość",
                    "Rodzaj towaru",
                    "Kilometry (telematyka)",
                    "Identyfikator wewnętrzny",
                ]
            )
            sheet.append(["ABC 123", "2026-06-01 06:34:31", 10.5, "Benzyna", 1000, 1])
            sheet.append(["???", "2026-06-02 06:34:31", 20, "Diesel", 0, 2])
            workbook.save(source)

            records = load_single_report(source)

        self.assertEqual(len(records), 2)
        self.assertEqual(records[0].vehicle, "ABC123")
        self.assertEqual(records[0].transaction_date, date(2026, 6, 1))
        self.assertEqual(records[0].product, "Benzyna")
        self.assertEqual(records[1].vehicle, UNKNOWN_VEHICLE)

    def test_analysis_counts_gasoline_as_fuel(self):
        records = [
            FuelRecord(SOURCE_TELEMATICS, "ABC123", date(2026, 6, 1), 10, 0, "EUR", "Benzyna", odometer=1000),
            FuelRecord(SOURCE_TELEMATICS, "ABC123", date(2026, 6, 2), 20, 0, "EUR", "Diesel", odometer=1100),
            FuelRecord(SOURCE_TELEMATICS, "ABC123", date(2026, 6, 3), 3, 0, "EUR", "AdBlue", odometer=1150),
        ]

        result = analyze_records(records, minimum_distance=1)
        row = result.rows[0]

        self.assertEqual(row.gasoline_liters, 10)
        self.assertEqual(row.diesel_total, 20)
        self.assertEqual(row.fuel_total, 30)
        self.assertEqual(row.adblue_liters, 3)
        self.assertAlmostEqual(row.consumption or 0, 20.0)
        self.assertEqual(result.total_gasoline, 10)
        self.assertEqual(result.total_fuel, 30)

    def test_analysis_builds_ranking_and_separates_adblue(self):
        result = analyze_records(self.records, minimum_distance=100)
        self.assertGreater(len(result.rows), 90)
        self.assertGreater(result.total_diesel, 0)
        self.assertGreater(result.total_adblue, 0)
        self.assertGreater(result.total_distance, 0)
        self.assertIsNotNone(result.worst)
        self.assertEqual(result.worst.rank, 1)

    def test_driver_override_groups_selected_vehicle(self):
        result = analyze_records(
            self.records,
            driver_overrides={"EL8HN73": "Jan Testowy"},
            minimum_distance=100,
        )
        row = next(item for item in result.rows if "EL8HN73" in item.vehicles)
        self.assertEqual(row.driver, "Jan Testowy")

    def test_archive_round_trip(self):
        with TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / "archiwum" / "raport.json"
            imported_at = datetime(2026, 6, 24, 12, 0, 0)
            save_archive_report(
                archive_path,
                self.records[:5],
                REPORT.name,
                imported_at=imported_at,
            )

            entries = list_archive_reports(archive_path.parent)
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0].source_file, REPORT.name)
            self.assertEqual(entries[0].record_count, 5)

            metadata, records = load_archive_report(archive_path)
            self.assertEqual(metadata["source_file"], REPORT.name)
            self.assertEqual(len(records), 5)
            self.assertEqual(records[0].vehicle, self.records[0].vehicle)
            self.assertEqual(records[0].transaction_date, self.records[0].transaction_date)

    def test_database_config_round_trip(self):
        with TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "baza.json"
            save_database_config(
                config_path,
                DatabaseConfig(
                    host="db.example.com",
                    port=5433,
                    database="fuel",
                    username="fuel_user",
                    password="secret",
                    sslmode="prefer",
                ),
            )

            raw_config = config_path.read_text(encoding="utf-8")
            if os.name == "nt":
                self.assertNotIn("secret", raw_config)

            config = load_database_config(config_path)
            self.assertTrue(config.is_complete)
            self.assertEqual(config.host, "db.example.com")
            self.assertEqual(config.port, 5433)
            self.assertEqual(config.database, "fuel")
            self.assertEqual(config.username, "fuel_user")
            self.assertEqual(config.password, "secret")

    def test_exports_valid_workbook(self):
        result = analyze_records(self.records, minimum_distance=100)
        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.xlsx"
            export_xlsx(result, output)
            self.assertTrue(output.exists())
            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Podsumowanie", "Transakcje", "Mapowanie kierowców"],
            )
            headers = [cell.value for cell in workbook["Podsumowanie"][11]]
            self.assertEqual(headers[4], "Diesel [l]")
            self.assertEqual(headers[5], "Benzyna [l]")
            self.assertEqual(headers[6], "Paliwo [l]")
            self.assertEqual(headers[7], "AdBlue [l]")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
