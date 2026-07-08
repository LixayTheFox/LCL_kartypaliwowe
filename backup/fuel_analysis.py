from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_TELEMATICS = "Raport paliwowy"
ARCHIVE_VERSION = 1


@dataclass(slots=True)
class FuelRecord:
    source: str
    vehicle: str
    transaction_date: date | None
    liters: float
    amount: float
    currency: str
    product: str
    driver: str = ""
    odometer: float | None = None
    provider: str = ""
    station: str = ""
    card_number: str = ""

    @property
    def is_adblue(self) -> bool:
        return "ADBLUE" in self.product.upper().replace(" ", "")

    @property
    def is_diesel(self) -> bool:
        normalized = self.product.upper().strip()
        return normalized in {"DIESEL", "ON", "OLEJ NAPĘDOWY", "OLEJ NAPEDOWY"}


@dataclass(slots=True)
class DriverSummary:
    driver: str
    vehicles: tuple[str, ...]
    transactions: int
    diesel_telematics: float
    diesel_cards: float
    diesel_total: float
    adblue_liters: float
    distance_km: float
    consumption: float | None
    cost_eur: float
    cost_pln: float
    other_costs: dict[str, float]
    first_date: date | None
    last_date: date | None
    rank: int | None = None
    status: str = ""

    @property
    def vehicle_label(self) -> str:
        return ", ".join(self.vehicles)


@dataclass(slots=True)
class AnalysisResult:
    rows: list[DriverSummary]
    records: list[FuelRecord]
    total_diesel: float
    total_adblue: float
    total_distance: float
    average_consumption: float | None
    ranked_count: int
    worst: DriverSummary | None
    date_from: date | None
    date_to: date | None
    warnings: list[str] = field(default_factory=list)


@dataclass(slots=True)
class ArchivedReportInfo:
    path: Path
    source_file: str
    imported_at: datetime | None
    date_from: date | None
    date_to: date | None
    record_count: int


def normalize_plate(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _as_float(value: object) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _as_optional_float(value: object) -> float | None:
    if value in (None, ""):
        return None
    number = _as_float(value)
    return number if number >= 0 else None


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _header_map(row: Iterable[object]) -> dict[str, int]:
    return {
        str(value).strip().lower(): index
        for index, value in enumerate(row)
        if value not in (None, "")
    }


def _column(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        index = headers.get(name.lower())
        if index is not None:
            return index
    return None


def _cell(row: tuple[object, ...], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def load_telematics(path: str | Path) -> list[FuelRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = _header_map(next(rows))
    except StopIteration as exc:
        raise ValueError("Raport XLSX jest pusty.") from exc

    vehicle_col = _column(headers, "Numery rejestracyjne", "Numer rejestracyjny", "Pojazd")
    date_col = _column(headers, "Data")
    liters_col = _column(headers, "Ilość", "Ilosc")
    product_col = _column(headers, "Rodzaj towaru", "Produkt", "Usługa")
    amount_col = _column(headers, "Kwota [EUR]", "Kwota EUR", "Kwota")
    odometer_col = _column(headers, "Kilometry (telematyka)", "Przebieg", "Kilometry")
    provider_col = _column(headers, "Dostawca kart", "Dostawca")
    station_col = _column(headers, "Stacja paliw", "Stacja")
    driver_col = _column(headers, "Kierowca")

    required = {
        "numer rejestracyjny": vehicle_col,
        "data": date_col,
        "ilość": liters_col,
        "rodzaj towaru": product_col,
    }
    missing = [name for name, index in required.items() if index is None]
    if missing:
        raise ValueError(
            "Raport nie ma wymaganych kolumn: " + ", ".join(missing)
        )

    records: list[FuelRecord] = []
    for row in rows:
        vehicle = normalize_plate(_cell(row, vehicle_col))
        transaction_date = _as_date(_cell(row, date_col))
        if not vehicle or transaction_date is None:
            continue
        records.append(
            FuelRecord(
                source=SOURCE_TELEMATICS,
                vehicle=vehicle,
                transaction_date=transaction_date,
                liters=_as_float(_cell(row, liters_col)),
                amount=_as_float(_cell(row, amount_col)),
                currency="EUR",
                product=str(_cell(row, product_col) or "").strip(),
                driver=str(_cell(row, driver_col) or "").strip(),
                odometer=_as_optional_float(_cell(row, odometer_col)),
                provider=str(_cell(row, provider_col) or "").strip(),
                station=str(_cell(row, station_col) or "").strip(),
            )
        )
    workbook.close()
    if not records:
        raise ValueError("W raporcie nie znaleziono żadnych transakcji.")
    return records


def load_single_report(path: str | Path) -> list[FuelRecord]:
    return load_telematics(path)


def analyze_records(
    records: Iterable[FuelRecord],
    driver_overrides: dict[str, str] | None = None,
    minimum_distance: float = 100.0,
) -> AnalysisResult:
    all_records = list(records)
    overrides = {
        normalize_plate(vehicle): driver.strip()
        for vehicle, driver in (driver_overrides or {}).items()
        if normalize_plate(vehicle) and driver.strip()
    }

    records_by_vehicle: dict[str, list[FuelRecord]] = defaultdict(list)
    for record in all_records:
        records_by_vehicle[record.vehicle].append(record)

    vehicle_driver: dict[str, str] = {}
    for vehicle, vehicle_records in records_by_vehicle.items():
        if vehicle in overrides:
            vehicle_driver[vehicle] = overrides[vehicle]
            continue
        candidates = [record.driver.strip() for record in vehicle_records if record.driver.strip()]
        vehicle_driver[vehicle] = Counter(candidates).most_common(1)[0][0] if candidates else ""

    grouped: dict[str, list[FuelRecord]] = defaultdict(list)
    group_vehicles: dict[str, set[str]] = defaultdict(set)
    for vehicle, vehicle_records in records_by_vehicle.items():
        driver = vehicle_driver[vehicle]
        group_key = driver.casefold() if driver else f"__vehicle__{vehicle}"
        grouped[group_key].extend(vehicle_records)
        group_vehicles[group_key].add(vehicle)

    summaries: list[DriverSummary] = []
    all_dates = [
        record.transaction_date for record in all_records if record.transaction_date is not None
    ]

    for group_key, group_records in grouped.items():
        vehicles = tuple(sorted(group_vehicles[group_key]))
        driver = vehicle_driver[vehicles[0]] if not group_key.startswith("__vehicle__") else ""
        driver_label = driver or f"Nieprzypisany - {vehicles[0]}"

        diesel_telematics = sum(
            record.liters for record in group_records if record.is_diesel
        )
        diesel_cards = 0.0
        adblue_liters = sum(record.liters for record in group_records if record.is_adblue)

        distance = 0.0
        for vehicle in vehicles:
            odometers = [
                record.odometer
                for record in records_by_vehicle[vehicle]
                if record.odometer is not None and record.odometer > 0
            ]
            if len(odometers) >= 2:
                distance += max(odometers) - min(odometers)

        diesel_total = diesel_telematics + diesel_cards
        consumption = diesel_total / distance * 100 if distance > 0 else None

        costs: dict[str, float] = defaultdict(float)
        for record in group_records:
            costs[record.currency or "EUR"] += record.amount

        dates = [
            record.transaction_date
            for record in group_records
            if record.transaction_date is not None
        ]
        if distance <= 0:
            status = "Brak dystansu"
        elif distance < minimum_distance:
            status = "Za krótki dystans"
        else:
            status = "W rankingu"

        summaries.append(
            DriverSummary(
                driver=driver_label,
                vehicles=vehicles,
                transactions=len(group_records),
                diesel_telematics=diesel_telematics,
                diesel_cards=diesel_cards,
                diesel_total=diesel_total,
                adblue_liters=adblue_liters,
                distance_km=distance,
                consumption=consumption,
                cost_eur=costs.pop("EUR", 0.0),
                cost_pln=costs.pop("PLN", 0.0),
                other_costs=dict(costs),
                first_date=min(dates) if dates else None,
                last_date=max(dates) if dates else None,
                status=status,
            )
        )

    eligible = sorted(
        (
            row
            for row in summaries
            if row.consumption is not None and row.distance_km >= minimum_distance
        ),
        key=lambda row: row.consumption or 0,
        reverse=True,
    )
    for rank, row in enumerate(eligible, start=1):
        row.rank = rank

    summaries.sort(
        key=lambda row: (
            row.rank is None,
            row.rank if row.rank is not None else 10**9,
            row.driver.casefold(),
        )
    )

    total_diesel = sum(row.diesel_total for row in summaries)
    total_adblue = sum(row.adblue_liters for row in summaries)
    total_distance = sum(row.distance_km for row in summaries)
    average_consumption = (
        total_diesel / total_distance * 100 if total_distance > 0 else None
    )

    warnings: list[str] = []
    if not any(record.driver.strip() for record in all_records):
        warnings.append(
            "Pliki nie zawierają nazwisk kierowców. Wyniki są pokazane per pojazd; "
            "nazwiska można przypisać w aplikacji."
        )
    if any(row.rank is None for row in summaries):
        warnings.append(
            "Część pozycji pominięto w rankingu z powodu braku lub zbyt małego dystansu."
        )

    return AnalysisResult(
        rows=summaries,
        records=all_records,
        total_diesel=total_diesel,
        total_adblue=total_adblue,
        total_distance=total_distance,
        average_consumption=average_consumption,
        ranked_count=len(eligible),
        worst=eligible[0] if eligible else None,
        date_from=min(all_dates) if all_dates else None,
        date_to=max(all_dates) if all_dates else None,
        warnings=warnings,
    )


def save_driver_mapping(path: str | Path, mapping: dict[str, str]) -> None:
    normalized = {
        normalize_plate(vehicle): driver.strip()
        for vehicle, driver in mapping.items()
        if normalize_plate(vehicle) and driver.strip()
    }
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def load_driver_mapping(path: str | Path) -> dict[str, str]:
    target = Path(path)
    if not target.exists():
        return {}
    try:
        raw = json.loads(target.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    if not isinstance(raw, dict):
        return {}
    return {
        normalize_plate(vehicle): str(driver).strip()
        for vehicle, driver in raw.items()
        if normalize_plate(vehicle) and str(driver).strip()
    }


def _date_to_iso(value: date | None) -> str | None:
    return value.isoformat() if value else None


def _datetime_to_iso(value: datetime | None) -> str | None:
    return value.isoformat(timespec="seconds") if value else None


def _date_from_iso(value: object) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return _as_date(value)


def _datetime_from_iso(value: object) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def _record_to_archive(record: FuelRecord) -> dict[str, object]:
    return {
        "source": record.source,
        "vehicle": record.vehicle,
        "transaction_date": _date_to_iso(record.transaction_date),
        "liters": record.liters,
        "amount": record.amount,
        "currency": record.currency,
        "product": record.product,
        "driver": record.driver,
        "odometer": record.odometer,
        "provider": record.provider,
        "station": record.station,
        "card_number": record.card_number,
    }


def _record_from_archive(raw: object) -> FuelRecord:
    if not isinstance(raw, dict):
        raise ValueError("Archiwum zawiera nieprawidłowy wpis transakcji.")
    return FuelRecord(
        source=str(raw.get("source") or SOURCE_TELEMATICS),
        vehicle=normalize_plate(raw.get("vehicle")),
        transaction_date=_date_from_iso(raw.get("transaction_date")),
        liters=_as_float(raw.get("liters")),
        amount=_as_float(raw.get("amount")),
        currency=str(raw.get("currency") or "EUR").strip().upper(),
        product=str(raw.get("product") or "").strip(),
        driver=str(raw.get("driver") or "").strip(),
        odometer=_as_optional_float(raw.get("odometer")),
        provider=str(raw.get("provider") or "").strip(),
        station=str(raw.get("station") or "").strip(),
        card_number=str(raw.get("card_number") or "").strip(),
    )


def save_archive_report(
    path: str | Path,
    records: Iterable[FuelRecord],
    source_file: str,
    imported_at: datetime | None = None,
) -> None:
    records_list = list(records)
    dates = [
        record.transaction_date
        for record in records_list
        if record.transaction_date is not None
    ]
    payload = {
        "version": ARCHIVE_VERSION,
        "source_file": source_file,
        "imported_at": _datetime_to_iso(imported_at or datetime.now()),
        "date_from": _date_to_iso(min(dates) if dates else None),
        "date_to": _date_to_iso(max(dates) if dates else None),
        "record_count": len(records_list),
        "records": [_record_to_archive(record) for record in records_list],
    }
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def load_archive_report(path: str | Path) -> tuple[dict[str, object], list[FuelRecord]]:
    target = Path(path)
    try:
        raw = json.loads(target.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError("Nie można odczytać archiwum raportu.") from exc
    if not isinstance(raw, dict):
        raise ValueError("Archiwum raportu ma nieprawidłowy format.")
    raw_records = raw.get("records")
    if not isinstance(raw_records, list):
        raise ValueError("Archiwum nie zawiera listy transakcji.")
    records = [_record_from_archive(item) for item in raw_records]
    if not records:
        raise ValueError("Archiwum nie zawiera transakcji.")
    return raw, records


def _archive_info_from_payload(path: Path, payload: dict[str, object]) -> ArchivedReportInfo:
    try:
        record_count = int(payload.get("record_count") or 0)
    except (TypeError, ValueError):
        record_count = 0
    if not record_count and isinstance(payload.get("records"), list):
        record_count = len(payload["records"])
    return ArchivedReportInfo(
        path=path,
        source_file=str(payload.get("source_file") or path.stem),
        imported_at=_datetime_from_iso(payload.get("imported_at")),
        date_from=_date_from_iso(payload.get("date_from")),
        date_to=_date_from_iso(payload.get("date_to")),
        record_count=record_count,
    )


def load_archive_info(path: str | Path) -> ArchivedReportInfo:
    target = Path(path)
    try:
        raw = json.loads(target.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError("Nie można odczytać archiwum raportu.") from exc
    if not isinstance(raw, dict):
        raise ValueError("Archiwum raportu ma nieprawidłowy format.")
    return _archive_info_from_payload(target, raw)


def list_archive_reports(directory: str | Path) -> list[ArchivedReportInfo]:
    target = Path(directory)
    if not target.exists():
        return []
    entries: list[ArchivedReportInfo] = []
    for path in target.glob("*.json"):
        try:
            entries.append(load_archive_info(path))
        except ValueError:
            continue
    return sorted(
        entries,
        key=lambda item: item.imported_at or datetime.min,
        reverse=True,
    )


def export_csv(result: AnalysisResult, path: str | Path) -> None:
    with Path(path).open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(
            [
                "Ranking",
                "Kierowca",
                "Pojazdy",
                "Liczba transakcji",
                "Diesel [l]",
                "AdBlue [l]",
                "Dystans [km]",
                "Spalanie [l/100 km]",
                "Koszt [EUR]",
                "Koszt [PLN]",
                "Status",
            ]
        )
        for row in result.rows:
            writer.writerow(
                [
                    row.rank or "",
                    row.driver,
                    row.vehicle_label,
                    row.transactions,
                    round(row.diesel_total, 2),
                    round(row.adblue_liters, 2),
                    round(row.distance_km, 1),
                    round(row.consumption, 2) if row.consumption is not None else "",
                    round(row.cost_eur, 2),
                    round(row.cost_pln, 2),
                    row.status,
                ]
            )


def export_xlsx(result: AnalysisResult, path: str | Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Podsumowanie"
    details = workbook.create_sheet("Transakcje")
    mapping = workbook.create_sheet("Mapowanie kierowców")

    navy = "14213D"
    blue = "2563EB"
    pale_blue = "EAF2FF"
    red = "C62828"
    pale_red = "FDECEC"
    pale_gray = "F3F4F6"
    green = "13795B"
    white = "FFFFFF"
    border = Border(bottom=Side(style="thin", color="D9E1EC"))

    summary.merge_cells("A1:L1")
    summary["A1"] = "Raport efektywności kierowców i pojazdów"
    summary["A1"].font = Font(size=20, bold=True, color=white)
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 36

    period = ""
    if result.date_from and result.date_to:
        period = f"{result.date_from:%d.%m.%Y} - {result.date_to:%d.%m.%Y}"
    summary["A3"] = "Okres"
    summary["B3"] = period
    summary["D3"] = "Wygenerowano"
    summary["E3"] = datetime.now().strftime("%d.%m.%Y %H:%M")

    cards = [
        ("A5", "Kierowcy / pojazdy", len(result.rows), "0"),
        ("D5", "Diesel razem", result.total_diesel, '#,##0.00 "l"'),
        ("G5", "Dystans", result.total_distance, '#,##0.0 "km"'),
        (
            "J5",
            "Średnie spalanie",
            result.average_consumption or 0,
            '0.00 "l/100 km"',
        ),
    ]
    for cell, label, value, number_format in cards:
        start_col = summary[cell].column
        summary.merge_cells(
            start_row=summary[cell].row,
            start_column=start_col,
            end_row=summary[cell].row,
            end_column=start_col + 2,
        )
        summary.merge_cells(
            start_row=summary[cell].row + 1,
            start_column=start_col,
            end_row=summary[cell].row + 2,
            end_column=start_col + 2,
        )
        label_cell = summary.cell(summary[cell].row, start_col)
        value_cell = summary.cell(summary[cell].row + 1, start_col)
        label_cell.value = label
        label_cell.font = Font(bold=True, color="536176")
        value_cell.value = value
        value_cell.font = Font(size=18, bold=True, color=navy)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(vertical="center")
        for row_cells in summary.iter_rows(
            min_row=summary[cell].row,
            max_row=summary[cell].row + 2,
            min_col=start_col,
            max_col=start_col + 2,
        ):
            for card_cell in row_cells:
                card_cell.fill = PatternFill("solid", fgColor=pale_blue)

    if result.worst:
        summary.merge_cells("A9:L9")
        summary["A9"] = (
            f"Najgorszy wynik: {result.worst.driver} | "
            f"{result.worst.consumption:.2f} l/100 km | "
            f"{result.worst.vehicle_label}"
        )
        summary["A9"].font = Font(bold=True, color=red)
        summary["A9"].fill = PatternFill("solid", fgColor=pale_red)
        summary["A9"].alignment = Alignment(vertical="center")
        summary.row_dimensions[9].height = 28

    headers = [
        "Ranking",
        "Kierowca",
        "Pojazdy",
        "Transakcje",
        "Diesel [l]",
        "AdBlue [l]",
        "Dystans [km]",
        "Spalanie [l/100 km]",
        "Koszt [EUR]",
        "Koszt [PLN]",
        "Status",
    ]
    header_row = 11
    summary.append([""] * len(headers))
    for col, header in enumerate(headers, start=1):
        cell = summary.cell(header_row, col, header)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[header_row].height = 34

    for row_index, row in enumerate(result.rows, start=header_row + 1):
        values = [
            row.rank,
            row.driver,
            row.vehicle_label,
            row.transactions,
            row.diesel_total,
            row.adblue_liters,
            row.distance_km,
            row.consumption,
            row.cost_eur,
            row.cost_pln,
            row.status,
        ]
        for col, value in enumerate(values, start=1):
            cell = summary.cell(row_index, col, value)
            cell.border = border
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
        if row.rank == 1:
            for cell in summary[row_index]:
                cell.fill = PatternFill("solid", fgColor=pale_red)
                cell.font = Font(bold=True, color=red)
        elif row.rank is None:
            for cell in summary[row_index]:
                cell.font = Font(color="7B8794")

    last_row = header_row + len(result.rows)
    if result.rows:
        table = Table(displayName="TabelaWynikow", ref=f"A{header_row}:K{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        summary.add_table(table)
        summary.conditional_formatting.add(
            f"H{header_row + 1}:H{last_row}",
            ColorScaleRule(
                start_type="min",
                start_color="C6EFCE",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB9C",
                end_type="max",
                end_color="FFC7CE",
            ),
        )

    for col in ("E", "F", "I", "J"):
        for cell in summary[col][header_row:]:
            cell.number_format = '#,##0.00'
    for cell in summary["G"][header_row:]:
        cell.number_format = '#,##0.0'
    for cell in summary["H"][header_row:]:
        cell.number_format = '0.00'

    widths = {
        "A": 10,
        "B": 29,
        "C": 25,
        "D": 11,
        "E": 14,
        "F": 13,
        "G": 15,
        "H": 21,
        "I": 14,
        "J": 14,
        "K": 19,
    }
    for col, width in widths.items():
        summary.column_dimensions[col].width = width
    summary.freeze_panes = f"A{header_row + 1}"
    summary.auto_filter.ref = f"A{header_row}:K{last_row}"
    summary.sheet_view.showGridLines = False

    detail_headers = [
        "Źródło",
        "Data",
        "Kierowca",
        "Pojazd",
        "Produkt",
        "Ilość [l]",
        "Kwota",
        "Waluta",
        "Przebieg [km]",
        "Dostawca",
        "Stacja",
        "Numer karty",
    ]
    details.append(detail_headers)
    vehicle_to_driver = {
        vehicle: row.driver for row in result.rows for vehicle in row.vehicles
    }
    for record in sorted(
        result.records,
        key=lambda item: (
            item.transaction_date or date.min,
            item.vehicle,
            item.source,
        ),
    ):
        details.append(
            [
                record.source,
                record.transaction_date,
                vehicle_to_driver.get(record.vehicle, record.driver),
                record.vehicle,
                record.product,
                record.liters,
                record.amount,
                record.currency,
                record.odometer,
                record.provider,
                record.station,
                record.card_number,
            ]
        )
    _style_data_sheet(details, "TabelaTransakcji", navy, white)
    details.column_dimensions["A"].width = 24
    details.column_dimensions["B"].width = 13
    details.column_dimensions["C"].width = 28
    details.column_dimensions["D"].width = 14
    details.column_dimensions["E"].width = 18
    details.column_dimensions["F"].width = 13
    details.column_dimensions["G"].width = 13
    details.column_dimensions["H"].width = 10
    details.column_dimensions["I"].width = 16
    details.column_dimensions["J"].width = 18
    details.column_dimensions["K"].width = 28
    details.column_dimensions["L"].width = 23
    for cell in details["B"][1:]:
        cell.number_format = "dd.mm.yyyy"
    for col in ("F", "G", "I"):
        for cell in details[col][1:]:
            cell.number_format = "#,##0.00"

    mapping.append(["Pojazd", "Kierowca"])
    for row in result.rows:
        for vehicle in row.vehicles:
            driver = "" if row.driver.startswith("Nieprzypisany -") else row.driver
            mapping.append([vehicle, driver])
    _style_data_sheet(mapping, "TabelaMapowania", green, white)
    mapping.column_dimensions["A"].width = 18
    mapping.column_dimensions["B"].width = 32

    workbook.save(path)


def _style_data_sheet(sheet, table_name: str, header_color: str, font_color: str) -> None:
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(bold=True, color=font_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    if sheet.max_row >= 2:
        ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
