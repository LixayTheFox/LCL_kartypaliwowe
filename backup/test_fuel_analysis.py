from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from fuel_analysis import (
    SOURCE_TELEMATICS,
    analyze_records,
    export_xlsx,
    list_archive_reports,
    load_archive_report,
    load_single_report,
    normalize_plate,
    save_archive_report,
)


BASE_DIR = Path(__file__).resolve().parent
REPORT = BASE_DIR / "export_20260601085249.xlsx"


class FuelAnalysisTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.records = load_single_report(REPORT)

    def test_normalizes_vehicle_registration(self):
        self.assertEqual(normalize_plate("BI 228HA"), "BI228HA")
        self.assertEqual(normalize_plate("bi-228 ha"), "BI228HA")

    def test_loads_single_report_format(self):
        self.assertEqual(len(self.records), 685)
        self.assertTrue(all(record.source == SOURCE_TELEMATICS for record in self.records))

    def test_analysis_builds_ranking_and_separates_adblue(self):
        result = analyze_records(self.records, minimum_distance=100)
        self.assertGreater(len(result.rows), 90)
        self.assertGreater(result.total_diesel, 0)
        self.assertGreater(result.total_adblue, 0)
        self.assertGreater(result.total_distance, 0)
        self.assertIsNotNone(result.worst)
        self.assertEqual(result.worst.rank, 1)

    def test_driver_override_groups_selected_vehicle(self):
        result = analyze_records(
            self.records,
            driver_overrides={"EL8HN73": "Jan Testowy"},
            minimum_distance=100,
        )
        row = next(item for item in result.rows if "EL8HN73" in item.vehicles)
        self.assertEqual(row.driver, "Jan Testowy")

    def test_archive_round_trip(self):
        with TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / "archiwum" / "raport.json"
            imported_at = datetime(2026, 6, 24, 12, 0, 0)
            save_archive_report(
                archive_path,
                self.records[:5],
                REPORT.name,
                imported_at=imported_at,
            )

            entries = list_archive_reports(archive_path.parent)
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0].source_file, REPORT.name)
            self.assertEqual(entries[0].record_count, 5)

            metadata, records = load_archive_report(archive_path)
            self.assertEqual(metadata["source_file"], REPORT.name)
            self.assertEqual(len(records), 5)
            self.assertEqual(records[0].vehicle, self.records[0].vehicle)
            self.assertEqual(records[0].transaction_date, self.records[0].transaction_date)

    def test_exports_valid_workbook(self):
        result = analyze_records(self.records, minimum_distance=100)
        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.xlsx"
            export_xlsx(result, output)
            self.assertTrue(output.exists())
            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Podsumowanie", "Transakcje", "Mapowanie kierowców"],
            )
            headers = [cell.value for cell in workbook["Podsumowanie"][11]]
            self.assertEqual(headers[4], "Diesel [l]")
            self.assertEqual(headers[5], "AdBlue [l]")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
